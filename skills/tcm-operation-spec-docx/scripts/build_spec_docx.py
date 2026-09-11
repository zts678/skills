# -*- coding: utf-8 -*-
"""
中医操作技术规范 docx 可复用生成器（tcm-operation-spec-docx 技能组件）

用法：
  1. 复制本脚本为工作目录下临时文件 gen_<技术名>.py；
  2. 只修改下方"内容填写区"的 meta 与 BLOCKS（排版函数禁止改动）；
  3. python gen_<技术名>.py
  4. 用 verify_layout.py 与既有基准文档比对；
  5. 交付后删除临时脚本。

内容块类型：
  ('h1', 文本) ('h2', 文本) ('h3', 文本)
  ('body', 文本)      正文/编号条目/参考文献条目（宋体12pt，1.5倍行距，首行缩进2字符）
  ('note', 文本)      灰色楷体注释（"注：……"）

自检：python build_spec_docx.py --selftest   （在系统临时目录生成样例 docx）
"""
import sys
import tempfile
from docx import Document
from docx.shared import Pt, Emu, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ==================== 排版层（禁止改动，参数见 references/format-spec.md） ====================

def _set_run(run, font, size, bold=False, color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rF = rPr.get_or_add_rFonts()
    rF.set(qn('w:ascii'), font)
    rF.set(qn('w:hAnsi'), font)
    rF.set(qn('w:eastAsia'), font)


def _new_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Emu(7772400)
    sec.page_height = Emu(10058400)
    sec.top_margin = Emu(914400)
    sec.bottom_margin = Emu(914400)
    sec.left_margin = Emu(1143000)
    sec.right_margin = Emu(1143000)
    st = doc.styles['Normal']
    st.font.name = '宋体'
    st.font.size = Pt(12)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    rf.set(qn('w:ascii'), '宋体')
    rf.set(qn('w:hAnsi'), '宋体')
    rf.set(qn('w:eastAsia'), '宋体')
    return doc


def _para(doc, line15=False, indent=False, align=None, before=None, after=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if align is not None:
        p.alignment = align
    if line15:
        pf.line_spacing = 1.5
        pf.space_after = Pt(3)
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if indent:
        pPr = p._p.get_or_add_pPr()
        ind = pPr.find(qn('w:ind'))
        if ind is None:
            ind = OxmlElement('w:ind')
            pPr.append(ind)
        ind.set(qn('w:firstLine'), '482')
    return p


def _cover(doc, meta):
    p = _para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    _set_run(p.add_run(meta['title']), '黑体', 22, bold=True)

    p = _para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
    r = p.add_run(meta['en1'])
    _set_run(r, '楷体', 14)
    r.add_break(WD_BREAK.LINE)
    _set_run(p.add_run(meta['en2']), '楷体', 14)

    p = _para(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_run(p.add_run('（操作技术规范）'), '楷体', 13)

    _para(doc)
    _para(doc)

    p = _para(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_run(p.add_run(meta['apply']), '宋体', 12)

    p = _para(doc)
    p.add_run().add_break(WD_BREAK.PAGE)


def _block(doc, kind, text):
    if kind == 'h1':
        p = _para(doc, before=12, after=6)
        _set_run(p.add_run(text), '黑体', 15, bold=True)
    elif kind == 'h2':
        p = _para(doc, before=8, after=4)
        _set_run(p.add_run(text), '黑体', 13, bold=True)
    elif kind == 'h3':
        p = _para(doc, before=6, after=3)
        _set_run(p.add_run(text), '黑体', 12, bold=True)
    elif kind == 'body':
        p = _para(doc, line15=True, indent=True)
        _set_run(p.add_run(text), '宋体', 12)
    elif kind == 'note':
        p = _para(doc, after=10)
        _set_run(p.add_run(text), '楷体', 10.5, color=RGBColor(0x59, 0x59, 0x59))
    else:
        raise ValueError('unknown block kind: %r' % kind)


def generate(meta, blocks, out_path):
    """meta: dict(title, en1, en2, apply)；blocks: [(kind, text), ...]"""
    doc = _new_doc()
    _cover(doc, meta)
    for item in blocks:
        _block(doc, item[0], item[1])
    doc.save(out_path)
    return len(doc.paragraphs)


# ==================== 内容填写区（复制后替换本区域及 main 输出路径） ====================

META = {
    'title': '中医示例技术操作技术规范',
    'en1': 'Technical operation specifications of example',
    'en2': 'technique in Chinese medicine',
    'apply': '适用于各级各类医疗机构相关专业执业医师',
}

BLOCKS = [
    ('h1', '编制说明'),
    ('body', '此处填写技术源流、编制所依据的标准/教材全称及标准号。'),
    ('note', '注：本文件供临床操作参考使用，具体执行应结合所在医疗机构规定及医嘱。'),
    ('h1', '1  范围'),
    ('body', '本规范规定了示例技术的术语和定义、操作步骤与要求、注意事项与禁忌。'),
    ('h1', '2  规范性引用文件'),
    ('body', '——GB/T 0000 示例标准'),
    ('h1', '3  术语和定义'),
    ('h2', '3.1  示例技术'),
    ('body', '示例技术的定义正文，应逐字来自权威资料原文。'),
    ('h3', '3.1.1  示例分型'),
    ('body', '示例分型的正文说明。'),
    ('h1', '9  参考文献'),
    ('body', '[1] 示例学会. 示例标准: T/X 0000—2026[S]. 示例地: 示例出版社, 2026.'),
    ('h1', '附录A  操作流程概要（资料性附录）'),
    ('body', '评估 → 准备 → 操作 → 术后处理 → 随访。'),
]

if __name__ == '__main__':
    if '--selftest' in sys.argv:
        out = tempfile.gettempdir() + r'\_tcm_skill_selftest.docx'
        n = generate(META, BLOCKS, out)
        print('selftest OK:', out, 'paragraphs:', n)
    else:
        print('请复制本脚本并替换 META/BLOCKS 后，或在代码中调用 generate(meta, blocks, out_path)。')
        print('自检：python build_spec_docx.py --selftest')
